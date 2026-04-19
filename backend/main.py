from fastapi import FastAPI, Request, HTTPException, UploadFile, File
from fastapi.responses import Response, FileResponse, StreamingResponse
from pydantic import BaseModel
import httpx
import os
import io
import json
import base64
import logging
import uuid
import jwt
from urllib.parse import quote
from datetime import datetime, timezone
from dotenv import load_dotenv
from dss_client import sign_document, get_access_token, get_certificates, verify_signature

logger = logging.getLogger("rb-office")
logging.basicConfig(level=logging.INFO)

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.platypus.flowables import Flowable
from pypdf import PdfReader, PdfWriter, Transformation
import pdfplumber

from docx import Document as DocxDocument
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import re

load_dotenv()

app = FastAPI(title="RB-Office Sign API")

JWT_SECRET = os.getenv("JWT_SECRET")
ONLYOFFICE_URL = os.getenv("ONLYOFFICE_URL")
ONLYOFFICE_INTERNAL_URL = os.getenv("ONLYOFFICE_INTERNAL_URL", "http://onlyoffice")
BACKEND_INTERNAL_URL = os.getenv("BACKEND_INTERNAL_URL", "http://rb-office-backend:8000")
SIGNATURES_DIR = os.getenv("SIGNATURES_DIR", "./signatures")
DOCUMENTS_DIR = os.getenv("DOCUMENTS_DIR", "./documents")

# Шрифт с поддержкой кириллицы — устанавливается через пакет fonts-dejavu-core
DEJAVU_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_BOLD_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
DEJAVU_MONO_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
_font_registered = False


def _register_fonts():
    global _font_registered
    if _font_registered:
        return
    if os.path.isfile(DEJAVU_PATH):
        pdfmetrics.registerFont(TTFont("DejaVu", DEJAVU_PATH))
    if os.path.isfile(DEJAVU_BOLD_PATH):
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", DEJAVU_BOLD_PATH))
    if os.path.isfile(DEJAVU_MONO_PATH):
        pdfmetrics.registerFont(TTFont("DejaVu-Mono", DEJAVU_MONO_PATH))
    _font_registered = True

os.makedirs(SIGNATURES_DIR, exist_ok=True)
os.makedirs(DOCUMENTS_DIR, exist_ok=True)

METADATA_FILE = os.path.join(DOCUMENTS_DIR, "metadata.json")


def _read_metadata() -> list:
    if not os.path.isfile(METADATA_FILE):
        return []
    with open(METADATA_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def _write_metadata(data: list):
    with open(METADATA_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _find_doc(docs: list, doc_id: str) -> dict | None:
    for d in docs:
        if d["id"] == doc_id:
            return d
    return None


def _get_document_type(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    if ext in ("doc", "docx", "odt", "rtf", "txt", "html", "htm", "pdf"):
        return "word"
    if ext in ("xls", "xlsx", "ods", "csv"):
        return "cell"
    if ext in ("ppt", "pptx", "odp"):
        return "slide"
    return "word"


# ── Health ──────────────────────────────────────────────────────

@app.get("/health")
async def health():
    return {"status": "ok"}


@app.get("/")
async def root():
    return {"message": "RB-Office Sign API is running"}


# ── DSS endpoints ──────────────────────────────────────────────

@app.get("/dss/check")
async def dss_check():
    """Проверяет подключение к DSS: получает токен и список сертификатов."""
    try:
        token = await get_access_token()
        certs = await get_certificates(token)
        return {
            "status": "ok",
            "token_received": True,
            "certificates": [
                {
                    "id": c.get("Id"),
                    "subject": c.get("SubjectName"),
                    "is_default": c.get("IsDefault"),
                    "valid_to": c.get("ValidTo"),
                }
                for c in certs
            ],
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/dss/sign-test")
async def sign_test():
    """Тестовое подписание — подписывает строку 'Hello DSS'."""
    try:
        test_content = b"Hello DSS - test signature"
        sig_bytes = await sign_document(test_content, "test.txt")
        sig_b64 = base64.b64encode(sig_bytes).decode()
        return {
            "status": "ok",
            "signature_size": len(sig_bytes),
            "signature_b64": sig_b64[:100] + "...",
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/dss/sign")
async def sign_endpoint(file: UploadFile = File(...)):
    """Принимает файл, возвращает отделённую подпись (.sig)."""
    try:
        file_content = await file.read()
        sig_bytes = await sign_document(file_content, file.filename)
        return Response(
            content=sig_bytes,
            media_type="application/octet-stream",
            headers={
                "Content-Disposition": f'attachment; filename="{file.filename}.sig"'
            },
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/dss/signatures/{key}")
async def get_signature(key: str):
    """Отдаёт .sig файл по ключу документа."""
    sig_path = os.path.join(SIGNATURES_DIR, f"{key}.sig")
    if not os.path.isfile(sig_path):
        raise HTTPException(status_code=404, detail="Подпись не найдена")
    with open(sig_path, "rb") as f:
        sig_bytes = f.read()
    return Response(
        content=sig_bytes,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{key}.sig"'},
    )


@app.post("/dss/verify")
async def verify_endpoint(
    file: UploadFile = File(..., description="Исходный документ"),
    signature: UploadFile = File(..., description="Файл подписи (.sig)"),
):
    """Проверяет отделённую подпись (два файла: документ + .sig)."""
    try:
        file_content = await file.read()
        sig_content = await signature.read()
        result = await verify_signature(file_content, sig_content)
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── OnlyOffice callback ───────────────────────────────────────

@app.post("/callback")
async def onlyoffice_callback(request: Request):
    """
    Callback от OnlyOffice при событиях редактирования.
    Статус 2 = документ сохранён, все пользователи вышли — скачиваем и сохраняем.
    """
    body = await request.json()
    status = body.get("status")
    key = body.get("key", "unknown")
    print(f"[CALLBACK] status={status} key={key} body={json.dumps(body, ensure_ascii=False)}")

    if status == 2:
        download_url = body.get("url")
        if not download_url:
            print(f"[CALLBACK] ERROR: status=2 but no URL in body")
            return {"error": 0}

        # Extract doc_id from key (format: "{uuid}_{timestamp}")
        # UUID has 5 parts: xxxxxxxx-xxxx-xxxx-xxxx-xxxxxxxxxxxx
        parts = key.split("_")
        doc_id = "_".join(parts[:-1]) if len(parts) > 1 else key

        docs = _read_metadata()
        doc = _find_doc(docs, doc_id)
        if not doc:
            print(f"[CALLBACK] ERROR: document not found for doc_id={doc_id}")
            return {"error": 0}

        # Download updated document from OnlyOffice
        # URL may use internal hostname — try as-is first, then replace with container name
        file_content = None
        async with httpx.AsyncClient(timeout=30.0) as client:
            try:
                response = await client.get(download_url)
                response.raise_for_status()
                file_content = response.content
                print(f"[CALLBACK] Downloaded {len(file_content)} bytes from {download_url}")
            except Exception as e:
                print(f"[CALLBACK] Failed to download from {download_url}: {e}")
                # Try replacing hostname with OnlyOffice container name
                from urllib.parse import urlparse, urlunparse
                parsed = urlparse(download_url)
                alt_url = urlunparse(parsed._replace(netloc="onlyoffice"))
                try:
                    response = await client.get(alt_url)
                    response.raise_for_status()
                    file_content = response.content
                    print(f"[CALLBACK] Downloaded {len(file_content)} bytes from {alt_url}")
                except Exception as e2:
                    print(f"[CALLBACK] Also failed from {alt_url}: {e2}")

        if file_content:
            # Save updated file, overwriting the original
            file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
            with open(file_path, "wb") as f:
                f.write(file_content)
            print(f"[CALLBACK] Saved updated document to {file_path}")
        else:
            print(f"[CALLBACK] Could not download document, skipping save")

        return {"error": 0}

    if status == 6:
        print(f"[CALLBACK] Force save error for key={key}")

    return {"error": 0}


# ── Document management ───────────────────────────────────────

@app.post("/documents/upload")
async def upload_document(file: UploadFile = File(...)):
    """Загружает документ, сохраняет на диск, добавляет в metadata.json."""
    doc_id = str(uuid.uuid4())
    ext = os.path.splitext(file.filename or "doc")[1]
    stored_name = f"{doc_id}{ext}"
    file_path = os.path.join(DOCUMENTS_DIR, stored_name)

    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    now = datetime.now(timezone.utc).isoformat()
    docs = _read_metadata()
    docs.append({
        "id": doc_id,
        "filename": file.filename,
        "stored_name": stored_name,
        "uploaded_at": now,
        "signed": False,
        "signature_time": None,
        "signer": None,
    })
    _write_metadata(docs)

    return {"id": doc_id, "filename": file.filename, "uploaded_at": now}


@app.get("/documents")
async def list_documents():
    """Список всех документов."""
    return _read_metadata()


@app.get("/documents/{doc_id}")
async def get_document(doc_id: str):
    """Возвращает полную запись о документе из metadata.json."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    return doc


@app.get("/documents/{doc_id}/download")
async def download_document(doc_id: str):
    """Отдаёт файл документа."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")
    return FileResponse(
        file_path,
        filename=doc["filename"],
        media_type="application/octet-stream",
    )


@app.delete("/documents/{doc_id}")
async def delete_document(doc_id: str):
    """Удаляет документ и его подпись (если есть)."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if os.path.isfile(file_path):
        os.remove(file_path)

    sig_path = os.path.join(SIGNATURES_DIR, f"{doc_id}.sig")
    if os.path.isfile(sig_path):
        os.remove(sig_path)

    docs = [d for d in docs if d["id"] != doc_id]
    _write_metadata(docs)

    return {"status": "deleted", "id": doc_id}


@app.post("/documents/{doc_id}/sign")
async def sign_doc(doc_id: str):
    """Подписывает документ через DSS, сохраняет подпись, обновляет метаданные."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    with open(file_path, "rb") as f:
        file_content = f.read()

    sig_bytes = await sign_document(file_content, doc["filename"])

    sig_path = os.path.join(SIGNATURES_DIR, f"{doc_id}.sig")
    with open(sig_path, "wb") as f:
        f.write(sig_bytes)

    # Верифицируем чтобы получить информацию о подписанте
    verify_result = await verify_signature(file_content, sig_bytes)

    signer_info = verify_result.get("signer", {}) or {}
    doc["signed"] = True
    doc["signature_time"] = verify_result.get("signing_time")
    doc["signer"] = signer_info.get("subject")
    doc["signer_info"] = signer_info
    doc["signature_type"] = verify_result.get("signature_type") or "CAdES-BES"
    doc["signature_valid"] = bool(verify_result.get("valid"))
    doc["verify_message"] = verify_result.get("message")
    _write_metadata(docs)

    return {
        "status": "signed",
        "id": doc_id,
        "signature_size": len(sig_bytes),
        "signer": doc["signer"],
        "signature_time": doc["signature_time"],
    }


class LocalSignPayload(BaseModel):
    signature_base64: str


@app.get("/documents/{doc_id}/content-base64")
async def get_document_content_base64(doc_id: str):
    """Отдаёт содержимое документа в base64 — для клиентского подписания в браузере."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    with open(file_path, "rb") as f:
        file_content = f.read()

    return {
        "content": base64.b64encode(file_content).decode(),
        "filename": doc["filename"],
    }


@app.post("/documents/{doc_id}/sign-local")
async def sign_doc_local(doc_id: str, payload: LocalSignPayload):
    """
    Принимает подпись, созданную в браузере через КриптоПро ЭЦП Browser plug-in.
    Сохраняет .sig, верифицирует через SVS, обновляет метаданные.
    """
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    try:
        sig_bytes = base64.b64decode(payload.signature_base64)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Некорректный base64: {e}")

    with open(file_path, "rb") as f:
        file_content = f.read()

    sig_path = os.path.join(SIGNATURES_DIR, f"{doc_id}.sig")
    with open(sig_path, "wb") as f:
        f.write(sig_bytes)

    verify_result = await verify_signature(file_content, sig_bytes)

    signer_info = verify_result.get("signer", {}) or {}
    doc["signed"] = True
    doc["signature_time"] = verify_result.get("signing_time")
    doc["signer"] = signer_info.get("subject")
    doc["signer_info"] = signer_info
    doc["signature_type"] = verify_result.get("signature_type") or "CAdES-BES"
    doc["signature_valid"] = bool(verify_result.get("valid"))
    doc["verify_message"] = verify_result.get("message")
    _write_metadata(docs)

    return {
        "status": "signed",
        "id": doc_id,
        "signature_size": len(sig_bytes),
        "signer": doc["signer"],
        "signature_time": doc["signature_time"],
        "verify": verify_result,
    }


class _LogoCircle(Flowable):
    """Круг 9мм с белыми буквами «РБ» — логотип RB-Office в шапке штампа."""

    def __init__(self, size: float = 9 * mm, fill: str = "#042C53", text_color: str = "#E6F1FB"):
        Flowable.__init__(self)
        self.size = size
        self.fill = fill
        self.text_color = text_color

    def wrap(self, _w, _h):
        return self.size, self.size

    def draw(self):
        c = self.canv
        r = self.size / 2
        c.setFillColor(HexColor(self.fill))
        c.circle(r, r, r, fill=1, stroke=0)
        c.setFillColor(HexColor(self.text_color))
        c.setFont("DejaVu-Bold", 10)
        # вертикально центрируем: ascender ~= 0.75 * fontSize
        c.drawCentredString(r, r - 3.1, "РБ")


def _extract_cn(subject: str) -> str:
    """Достаёт значение CN= из subject в стиле RFC 2253 / openssl."""
    if not subject:
        return ""
    fields = _parse_subject(subject)
    if "CN" in fields:
        return fields["CN"].strip()
    m = re.search(r"(?:^|,\s*|/)CN=([^,/]+)", subject)
    return (m.group(1).strip() if m else subject).strip()


def _parse_subject(subject: str) -> dict:
    # Разбирает subject сертификата в стиле RFC 2253 — поля через запятую,
    # значения могут быть в двойных кавычках, внутри которых "" означает
    # одну литеральную кавычку. Ключи нормализуются в UPPERCASE.
    result: dict = {}
    if not subject:
        return result
    s = subject
    i = 0
    n = len(s)
    while i < n:
        while i < n and s[i] in " ,":
            i += 1
        if i >= n:
            break
        key_start = i
        while i < n and s[i] != "=":
            i += 1
        if i >= n:
            break
        key = s[key_start:i].strip().upper()
        i += 1  # пропускаем '='
        if i < n and s[i] == '"':
            i += 1
            buf_chars = []
            while i < n:
                if s[i] == '"':
                    if i + 1 < n and s[i + 1] == '"':
                        buf_chars.append('"')
                        i += 2
                    else:
                        i += 1
                        break
                else:
                    buf_chars.append(s[i])
                    i += 1
            value = "".join(buf_chars)
        else:
            val_start = i
            while i < n and s[i] != ",":
                i += 1
            value = s[val_start:i].strip()
        if key:
            result[key] = value
    return result


def _fmt_date_short(value: str) -> str:
    """Пытается привести произвольное представление даты к DD.MM.YYYY."""
    if not value:
        return ""
    s = str(value).strip()
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).strftime("%d.%m.%Y")
    except Exception:
        pass
    m = re.search(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", s)
    if m:
        d, mo, y = m.group(1), m.group(2), m.group(3)
        if len(y) == 2:
            y = "20" + y
        return f"{int(d):02d}.{int(mo):02d}.{y}"
    return s


def _fmt_sign_time(value: str) -> str:
    """«DD.MM.YYYY, HH:MM». Источники: ISO 8601 или «HH:MM DD.MM.YYYY (…)»."""
    if not value:
        return ""
    s = str(value).strip()
    try:
        dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return dt.strftime("%d.%m.%Y, %H:%M")
    except Exception:
        pass
    m = re.search(r"(\d{1,2}):(\d{2})\s+(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", s)
    if m:
        h, mi, d, mo, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        return f"{int(d):02d}.{int(mo):02d}.{y}, {int(h):02d}:{int(mi):02d}"
    return s


class _StatusBadge(Flowable):
    """Зелёный бейдж «✓ Подпись верна» с закруглёнными углами."""

    def __init__(
        self,
        text: str = "Подпись верна",
        fill: str = "#EAF3DE",
        text_color: str = "#27500A",
        width: float = 33 * mm,
        height: float = 5.5 * mm,
    ):
        Flowable.__init__(self)
        self.text = text
        self.fill = fill
        self.text_color = text_color
        self.w = width
        self.h = height

    def wrap(self, _aw, _ah):
        return self.w, self.h

    def draw(self):
        c = self.canv
        c.setFillColor(HexColor(self.fill))
        c.setStrokeColor(HexColor(self.fill))
        c.roundRect(0, 0, self.w, self.h, 1.8 * mm, fill=1, stroke=0)
        c.setFillColor(HexColor(self.text_color))
        c.setFont("DejaVu-Bold", 7)
        c.drawString(2.2 * mm, self.h / 2 - 2.3, "✓ " + self.text)


def _build_stamp_flowables(doc: dict, avail_w: float) -> list:
    """Возвращает список flowables штампа (шапка + таблица подписей) на заданную ширину."""
    _register_fonts()

    registered = pdfmetrics.getRegisteredFontNames()
    font_regular = "DejaVu" if "DejaVu" in registered else "Helvetica"
    font_bold = "DejaVu-Bold" if "DejaVu-Bold" in registered else "Helvetica-Bold"
    font_mono = "DejaVu-Mono" if "DejaVu-Mono" in registered else "Courier"

    signer_info = doc.get("signer_info") or {}
    is_valid = bool(doc.get("signature_valid", True))
    subject_raw = signer_info.get("subject") or doc.get("signer") or ""
    subject_fields = _parse_subject(subject_raw)
    cn = _extract_cn(subject_raw) or subject_raw or "—"
    org_name = subject_fields.get("O", "").strip()
    position = subject_fields.get("T", "").strip()
    surname = subject_fields.get("SN", "").strip()
    given_name = subject_fields.get("G", "").strip()
    full_name = f"{surname} {given_name}".strip()

    # ── Стили абзацев для ячеек ─────────────────────────────────────
    st_header_title = ParagraphStyle(
        "h_title", fontName=font_regular, fontSize=11,
        textColor=HexColor("#1A1A1A"), leading=13,
    )
    st_header_sub = ParagraphStyle(
        "h_sub", fontName=font_regular, fontSize=8,
        textColor=HexColor("#5F5E5A"), leading=10,
    )
    st_col_header = ParagraphStyle(
        "col_h", fontName=font_regular, fontSize=8,
        textColor=HexColor("#5F5E5A"), leading=10,
    )
    st_role = ParagraphStyle(
        "role", fontName=font_bold, fontSize=9,
        textColor=HexColor("#0C447C"), leading=11,
    )
    st_org = ParagraphStyle(
        "org", fontName=font_bold, fontSize=9,
        textColor=HexColor("#1A1A1A"), leading=11,
    )
    st_org_sub = ParagraphStyle(
        "org_sub", fontName=font_regular, fontSize=8,
        textColor=HexColor("#5F5E5A"), leading=10,
    )
    st_cell = ParagraphStyle(
        "cell", fontName=font_regular, fontSize=9,
        textColor=HexColor("#1A1A1A"), leading=11,
    )
    st_mute = ParagraphStyle(
        "mute", fontName=font_regular, fontSize=9,
        textColor=HexColor("#5F5E5A"), leading=11,
    )
    st_mono = ParagraphStyle(
        "mono", fontName=font_mono, fontSize=7,
        textColor=HexColor("#1A1A1A"), leading=9,
    )
    st_small_mute = ParagraphStyle(
        "smute", fontName=font_regular, fontSize=8,
        textColor=HexColor("#5F5E5A"), leading=10,
    )

    # ── Шапка: лого + текст ─────────────────────────────────────────
    doc_id_text = doc.get("id") or ""
    header_right = [
        Paragraph("Документ подписан через RB-Office", st_header_title),
        Spacer(1, 1 * mm),
        Paragraph(f"Идентификатор {doc_id_text}", st_header_sub),
    ]
    header_table = Table(
        [[_LogoCircle(9 * mm), header_right]],
        colWidths=[12 * mm, avail_w - 12 * mm],
    )
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, HexColor("#D3D1C7")),
    ]))

    # ── Таблица подписей ───────────────────────────────────────────
    # Доли: РОЛЬ 12% / ОРГ 28% / ДОВЕРЕННОСТЬ 18% / СЕРТИФИКАТ 22% / ДАТА 20%
    col_widths = [avail_w * p for p in (0.12, 0.28, 0.18, 0.22, 0.20)]

    header_cells = [
        Paragraph("РОЛЬ", st_col_header),
        Paragraph("ОРГАНИЗАЦИЯ / СОТРУДНИК", st_col_header),
        Paragraph("ДОВЕРЕННОСТЬ", st_col_header),
        Paragraph("СЕРТИФИКАТ", st_col_header),
        Paragraph("ДАТА ПОДПИСАНИЯ", st_col_header),
    ]

    serial = signer_info.get("serial") or "—"
    valid_range = f"{_fmt_date_short(signer_info.get('valid_from')) or '—'} — {_fmt_date_short(signer_info.get('valid_to')) or '—'}"
    sign_time_str = _fmt_sign_time(doc.get("signature_time"))

    cert_cell = [
        Paragraph(serial, st_mono),
        Spacer(1, 1.5 * mm),
        Paragraph(valid_range, st_small_mute),
    ]

    date_cell = [
        Paragraph(sign_time_str or "—", st_cell),
        Spacer(1, 1.5 * mm),
        _StatusBadge("Подпись верна" if is_valid else "Подпись неверна",
                     fill="#EAF3DE" if is_valid else "#F7DCDC",
                     text_color="#27500A" if is_valid else "#7A1A1A"),
    ]

    st_position = ParagraphStyle(
        "position", fontName=font_regular, fontSize=8,
        textColor=HexColor("#5F5E5A"), leading=10, spaceBefore=2,
    )
    st_name = ParagraphStyle(
        "name", fontName=font_regular, fontSize=9,
        textColor=HexColor("#1A1A1A"), leading=11, spaceBefore=2,
    )

    org_cell: list = []
    if org_name:
        org_cell.append(Paragraph(org_name, st_org))
    if position:
        org_cell.append(Paragraph(position, st_position))
    if full_name:
        org_cell.append(Paragraph(full_name, st_name))
    if not org_cell:
        org_cell.append(Paragraph(cn or "—", st_org))

    signer_row = [
        Paragraph("Отправитель", st_role),
        org_cell,
        Paragraph("Не требуется", st_mute),
        cert_cell,
        date_cell,
    ]

    sig_table = Table(
        [header_cells, signer_row],
        colWidths=col_widths,
        repeatRows=1,
    )
    sig_table.setStyle(TableStyle([
        # Шапка колонок
        ("BACKGROUND", (0, 0), (-1, 0), HexColor("#F1EFE8")),
        ("TOPPADDING", (0, 0), (-1, 0), 2.5 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 2.5 * mm),
        ("LEFTPADDING", (0, 0), (-1, 0), 3 * mm),
        ("RIGHTPADDING", (0, 0), (-1, 0), 3 * mm),
        # Строки данных
        ("VALIGN", (0, 1), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 1), (-1, -1), 4 * mm),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4 * mm),
        ("LEFTPADDING", (0, 1), (-1, -1), 3 * mm),
        ("RIGHTPADDING", (0, 1), (-1, -1), 3 * mm),
        # Разделитель над строкой подписи
        ("LINEABOVE", (0, 1), (-1, 1), 0.5, HexColor("#D3D1C7")),
        ("LINEBELOW", (0, -1), (-1, -1), 0.5, HexColor("#D3D1C7")),
    ]))

    return [header_table, Spacer(1, 4 * mm), sig_table]


def _measure_flowables_height(flowables: list, avail_w: float) -> float:
    """Суммарная высота списка flowables при отрисовке на ширине avail_w (в точках)."""
    huge = 10_000 * mm
    total = 0.0
    for f in flowables:
        _w, h = f.wrap(avail_w, huge)
        total += h
    return total


def _build_stamp_standalone(doc: dict) -> bytes:
    """A4-страница штампа с полями 15 мм — используется когда штамп не влезает на последнюю страницу документа."""
    margin = 15 * mm
    avail_w = A4[0] - 2 * margin
    flowables = _build_stamp_flowables(doc, avail_w)
    buf = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=margin, rightMargin=margin,
        topMargin=margin, bottomMargin=margin,
        title="RB-Office: штамп подписи",
    )
    pdf_doc.build(flowables)
    return buf.getvalue()


def _build_stamp_overlay(doc: dict, avail_w: float) -> tuple:
    """
    PDF-страница размера (avail_w, высота-под-штамп) без полей — для оверлея
    на последнюю страницу документа. Возвращает (pdf_bytes, stamp_height_pt).
    """
    flowables = _build_stamp_flowables(doc, avail_w)
    stamp_h = _measure_flowables_height(flowables, avail_w)
    page_w = avail_w
    page_h = stamp_h

    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=(page_w, page_h))
    y = page_h
    for f in flowables:
        _w, h = f.wrap(avail_w, page_h)
        y -= h
        f.drawOn(c, 0, y)
    c.showPage()
    c.save()
    return buf.getvalue(), stamp_h


def _find_last_page_content_bottom_y(pdf_bytes: bytes) -> float:
    """
    Возвращает Y-координату нижнего края контента последней страницы в точках,
    отсчитывая от низа страницы (PDF-координаты снизу вверх).
    Если страница пустая — возвращает page.height (весь низ свободен).
    При ошибке парсинга — 0 (считаем что места нет, запасной вариант).
    """
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            page = pdf.pages[-1]
            page_h = float(page.height)
            chars = page.chars or []
            if not chars:
                return page_h
            # pdfplumber: 'bottom' — Y от верха страницы (top-down); переводим в PDF (bottom-up).
            max_bottom_td = max(c["bottom"] for c in chars)
            return page_h - max_bottom_td
    except Exception as e:
        logger.warning("[download-pdf] pdfplumber failed: %s — fallback to standalone stamp page", e)
        return 0.0


# ── DOCX stamp helpers ────────────────────────────────────────────

def _shade_cell(cell, hex_color: str):
    """Заливка ячейки через <w:shd> в tcPr."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _shade_run(run, hex_color: str):
    """Фон под буквами — <w:shd> в rPr (highlighting на уровне run)."""
    r_pr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    r_pr.append(shd)


def _set_table_borders(table, size: int, color: str):
    """Единый тонкий border по всем рёбрам таблицы. size в восьмых пункта (4 = 0.5pt)."""
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        tbl_borders.append(b)
    tbl_pr.append(tbl_borders)


def _set_cell_width(cell, width_mm: float):
    """Жёстко фиксирует ширину ячейки через <w:tcW> (autofit сам по себе не держит)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tcW = tc_pr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tc_pr.append(tcW)
    # 1 mm ≈ 56.69 twentieths of a point (twip); используем dxa (twips)
    tcW.set(qn("w:w"), str(int(width_mm * 56.7)))
    tcW.set(qn("w:type"), "dxa")


def _add_signature_stamp_to_docx(docx_bytes: bytes, doc_meta: dict) -> bytes:
    """
    Дорисовывает в конец DOCX заголовок «Документ подписан через RB-Office»,
    строку с идентификатором и таблицу подписей в стиле PDF-штампа.
    """
    doc = DocxDocument(io.BytesIO(docx_bytes))

    signer_info = doc_meta.get("signer_info") or {}
    subject_raw = signer_info.get("subject") or doc_meta.get("signer") or ""
    subject_fields = _parse_subject(subject_raw)
    org_name = subject_fields.get("O", "").strip()
    position = subject_fields.get("T", "").strip()
    surname = subject_fields.get("SN", "").strip()
    given_name = subject_fields.get("G", "").strip()
    full_name = f"{surname} {given_name}".strip()
    cn = _extract_cn(subject_raw)
    is_valid = bool(doc_meta.get("signature_valid", True))

    # Отступ от контента
    doc.add_paragraph()

    # Заголовок
    p_header = doc.add_paragraph()
    run_h = p_header.add_run("Документ подписан через RB-Office")
    run_h.font.size = Pt(11)
    run_h.font.bold = True
    run_h.font.color.rgb = RGBColor(0x04, 0x2C, 0x53)

    # ID
    p_id = doc.add_paragraph()
    run_id = p_id.add_run(f"Идентификатор {doc_meta.get('id') or ''}")
    run_id.font.size = Pt(8)
    run_id.font.color.rgb = RGBColor(0x5F, 0x5E, 0x5A)

    # Таблица: 5 колонок × 2 строки
    table = doc.add_table(rows=2, cols=5)
    table.autofit = False
    widths_mm = [20, 48, 30, 38, 34]  # сумма 170 — ширина текста A4 за вычетом полей 2×20 мм

    headers = ["РОЛЬ", "ОРГАНИЗАЦИЯ / СОТРУДНИК", "ДОВЕРЕННОСТЬ", "СЕРТИФИКАТ", "ДАТА ПОДПИСАНИЯ"]
    header_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = header_row.cells[i]
        _set_cell_width(cell, widths_mm[i])
        _shade_cell(cell, "F1EFE8")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.size = Pt(7.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x5F, 0x5E, 0x5A)

    data_row = table.rows[1]
    for i in range(5):
        _set_cell_width(data_row.cells[i], widths_mm[i])
        data_row.cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # РОЛЬ
    cell_role = data_row.cells[0]
    r = cell_role.paragraphs[0].add_run("Отправитель")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0C, 0x44, 0x7C)

    # ОРГ / СОТРУДНИК — до трёх строк
    cell_org = data_row.cells[1]
    cell_org.paragraphs[0].text = ""
    first_p = cell_org.paragraphs[0]
    if org_name:
        r1 = first_p.add_run(org_name)
        r1.font.size = Pt(9)
        r1.font.bold = True
    if position:
        p2 = first_p if not org_name else cell_org.add_paragraph()
        r2 = p2.add_run(position)
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor(0x5F, 0x5E, 0x5A)
    if full_name:
        p3 = first_p if not (org_name or position) else cell_org.add_paragraph()
        r3 = p3.add_run(full_name)
        r3.font.size = Pt(9)
    if not (org_name or position or full_name):
        rfb = first_p.add_run(cn or "—")
        rfb.font.size = Pt(9)
        rfb.font.bold = True

    # ДОВЕРЕННОСТЬ
    cell_proxy = data_row.cells[2]
    r = cell_proxy.paragraphs[0].add_run("Не требуется")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5F, 0x5E, 0x5A)

    # СЕРТИФИКАТ
    cell_cert = data_row.cells[3]
    cell_cert.paragraphs[0].text = ""
    p_serial = cell_cert.paragraphs[0]
    r_serial = p_serial.add_run(signer_info.get("serial") or "—")
    r_serial.font.size = Pt(7)
    r_serial.font.name = "Consolas"
    valid_range = (
        f"{_fmt_date_short(signer_info.get('valid_from')) or '—'} — "
        f"{_fmt_date_short(signer_info.get('valid_to')) or '—'}"
    )
    p_valid = cell_cert.add_paragraph()
    r_valid = p_valid.add_run(valid_range)
    r_valid.font.size = Pt(8)
    r_valid.font.color.rgb = RGBColor(0x5F, 0x5E, 0x5A)

    # ДАТА + бейдж
    cell_date = data_row.cells[4]
    cell_date.paragraphs[0].text = ""
    p_dt = cell_date.paragraphs[0]
    r_dt = p_dt.add_run(_fmt_sign_time(doc_meta.get("signature_time")) or "—")
    r_dt.font.size = Pt(9)
    p_badge = cell_date.add_paragraph()
    badge_text = " ✓ Подпись верна " if is_valid else " ✗ Подпись неверна "
    r_badge = p_badge.add_run(badge_text)
    r_badge.font.size = Pt(7)
    r_badge.font.bold = True
    if is_valid:
        r_badge.font.color.rgb = RGBColor(0x27, 0x50, 0x0A)
        _shade_run(r_badge, "EAF3DE")
    else:
        r_badge.font.color.rgb = RGBColor(0x7A, 0x1A, 0x1A)
        _shade_run(r_badge, "F7DCDC")

    _set_table_borders(table, size=4, color="D3D1C7")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def _convert_to_pdf(doc_id: str, filename: str, ext: str) -> bytes:
    """Запрашивает у OnlyOffice ConvertService конвертацию документа в PDF."""
    convert_key = f"{doc_id}_pdf_{int(datetime.now(timezone.utc).timestamp())}"

    payload = {
        "async": False,
        "filetype": ext.lower().lstrip("."),
        "outputtype": "pdf",
        "key": convert_key,
        "title": filename,
        "url": f"{BACKEND_INTERNAL_URL}/documents/{doc_id}/download",
    }
    payload["token"] = jwt.encode(payload, JWT_SECRET, algorithm="HS256")

    convert_url = f"{ONLYOFFICE_INTERNAL_URL}/ConvertService.ashx"

    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            convert_url,
            json=payload,
            headers={"Accept": "application/json"},
        )
        if resp.status_code != 200:
            raise HTTPException(
                status_code=502,
                detail=f"OnlyOffice ConvertService HTTP {resp.status_code}: {resp.text[:300]}",
            )
        try:
            result = resp.json()
        except Exception:
            raise HTTPException(
                status_code=502,
                detail=f"OnlyOffice вернул не-JSON: {resp.text[:300]}",
            )

        if result.get("error"):
            raise HTTPException(
                status_code=502,
                detail=f"OnlyOffice ConvertService error={result.get('error')}",
            )

        file_url = result.get("fileUrl")
        if not file_url:
            raise HTTPException(
                status_code=502,
                detail=f"ConvertService не вернул fileUrl: {result}",
            )

        pdf_resp = await client.get(file_url)
        pdf_resp.raise_for_status()
        return pdf_resp.content


# OnlyOffice ConvertService поддерживает конвертацию в PDF из office-форматов,
# но не из plain text. Явный белый список — чтобы давать пользователю понятную ошибку.
_PDF_CONVERTIBLE_EXTS = {
    "docx", "doc", "odt", "rtf",
    "xlsx", "xls", "ods", "csv",
    "pptx", "ppt", "odp",
    "pdf",
}


def _ascii_fallback(name: str) -> str:
    """Грубый ASCII-фоллбек для filename= в Content-Disposition."""
    return name.encode("ascii", "ignore").decode("ascii") or "document.pdf"


@app.get("/documents/{doc_id}/download-pdf")
async def download_pdf(doc_id: str):
    """Конвертирует документ в PDF и добавляет страницу-штамп с информацией о подписи."""
    logger.info("[download-pdf] start doc_id=%s", doc_id)

    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    if not doc.get("signed"):
        raise HTTPException(status_code=400, detail="Документ не подписан")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    ext = os.path.splitext(doc["filename"])[1].lstrip(".").lower()
    if ext not in _PDF_CONVERTIBLE_EXTS:
        raise HTTPException(
            status_code=400,
            detail=f"Конвертация в PDF не поддерживается для файлов .{ext}",
        )

    try:
        logger.info("[download-pdf] converting via OnlyOffice: filename=%r ext=%s", doc["filename"], ext)
        orig_pdf_bytes = await _convert_to_pdf(doc_id, doc["filename"], ext)
        logger.info("[download-pdf] converted: %d bytes", len(orig_pdf_bytes))

        content_reader = PdfReader(io.BytesIO(orig_pdf_bytes))
        last_page = content_reader.pages[-1]
        page_w_pt = float(last_page.mediabox.width)
        page_h_pt = float(last_page.mediabox.height)

        margin_pt = 15 * mm
        gap_pt = 5 * mm
        side_avail_w = page_w_pt - 2 * margin_pt

        overlay_bytes, stamp_h_pt = _build_stamp_overlay(doc, side_avail_w)
        content_bottom_y = _find_last_page_content_bottom_y(orig_pdf_bytes)
        needed = stamp_h_pt + gap_pt + margin_pt
        fits = content_bottom_y >= needed
        logger.info(
            "[download-pdf] last-page geometry: width=%.1fpt height=%.1fpt content_bottom_y=%.1fpt needed=%.1fpt fits=%s",
            page_w_pt, page_h_pt, content_bottom_y, needed, fits,
        )

        writer = PdfWriter()

        if fits:
            # Оверлеим штамп на последнюю страницу: низ штампа на margin_pt от низа.
            stamp_reader = PdfReader(io.BytesIO(overlay_bytes))
            stamp_page = stamp_reader.pages[0]
            tx = margin_pt
            ty = margin_pt
            last_page.merge_transformed_page(
                stamp_page, Transformation().translate(tx=tx, ty=ty)
            )
            for page in content_reader.pages:
                writer.add_page(page)
            logger.info("[download-pdf] stamp embedded on last page at ty=%.1fpt", ty)
        else:
            # Не влезло — добавляем отдельной A4-страницей.
            for page in content_reader.pages:
                writer.add_page(page)
            standalone_bytes = _build_stamp_standalone(doc)
            standalone_reader = PdfReader(io.BytesIO(standalone_bytes))
            for page in standalone_reader.pages:
                writer.add_page(page)
            logger.info("[download-pdf] stamp appended as new page")

        out = io.BytesIO()
        writer.write(out)
        out.seek(0)
        logger.info("[download-pdf] final size: %d bytes", len(out.getvalue()))
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("[download-pdf] failed for doc_id=%s: %s", doc_id, e)
        raise HTTPException(status_code=500, detail=f"Ошибка генерации PDF: {e}")

    base_name = os.path.splitext(doc["filename"])[0]
    out_name = f"{base_name}_signed.pdf"
    ascii_name = _ascii_fallback(out_name)
    # RFC 5987: filename= для ASCII-клиентов + filename*= для UTF-8.
    # HTTP-заголовки должны быть латин-1, поэтому кириллицу — только через %-кодирование.
    disposition = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(out_name)}"

    return StreamingResponse(
        out,
        media_type="application/pdf",
        headers={"Content-Disposition": disposition},
    )


@app.get("/documents/{doc_id}/download-docx-signed")
async def download_docx_signed(doc_id: str):
    """Скачивает исходный .docx с дорисованным в конце штампом подписи."""
    logger.info("[download-docx-signed] start doc_id=%s", doc_id)

    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    if not doc.get("signed"):
        raise HTTPException(status_code=400, detail="Документ не подписан")

    filename = doc.get("filename") or ""
    ext = os.path.splitext(filename)[1].lstrip(".").lower()
    if ext != "docx":
        raise HTTPException(
            status_code=400,
            detail=f"Вставка штампа в DOCX доступна только для .docx файлов (этот файл .{ext})",
        )

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    try:
        with open(file_path, "rb") as f:
            original_bytes = f.read()
        stamped_bytes = _add_signature_stamp_to_docx(original_bytes, doc)
        logger.info(
            "[download-docx-signed] stamped: original=%d bytes, out=%d bytes",
            len(original_bytes), len(stamped_bytes),
        )
    except Exception as e:
        logger.exception("[download-docx-signed] failed for doc_id=%s: %s", doc_id, e)
        raise HTTPException(status_code=500, detail=f"Ошибка генерации DOCX со штампом: {e}")

    base_name = os.path.splitext(filename)[0]
    out_name = f"{base_name}_signed.docx"
    ascii_name = _ascii_fallback(out_name)
    disposition = f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(out_name)}"

    return StreamingResponse(
        io.BytesIO(stamped_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": disposition},
    )


@app.post("/documents/{doc_id}/verify")
async def verify_doc(doc_id: str):
    """Проверяет подпись документа."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(DOCUMENTS_DIR, doc["stored_name"])
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на диске")

    sig_path = os.path.join(SIGNATURES_DIR, f"{doc_id}.sig")
    if not os.path.isfile(sig_path):
        raise HTTPException(status_code=404, detail="Подпись не найдена")

    with open(file_path, "rb") as f:
        file_content = f.read()
    with open(sig_path, "rb") as f:
        sig_content = f.read()

    result = await verify_signature(file_content, sig_content)
    return result


@app.get("/documents/{doc_id}/editor-config")
async def editor_config(doc_id: str):
    """Возвращает конфиг для OnlyOffice JS API, подписанный JWT."""
    docs = _read_metadata()
    doc = _find_doc(docs, doc_id)
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")

    ext = os.path.splitext(doc["filename"])[1].lstrip(".")
    editor_key = f"{doc_id}_{int(datetime.now(timezone.utc).timestamp())}"

    backend_url = BACKEND_INTERNAL_URL
    is_signed = bool(doc.get("signed"))

    document_cfg = {
        "fileType": ext,
        "key": editor_key,
        "title": doc["filename"],
        "url": f"{backend_url}/documents/{doc_id}/download",
    }
    if is_signed:
        document_cfg["permissions"] = {
            "edit": False,
            "download": True,
            "print": True,
            "review": False,
            "comment": False,
        }

    config = {
        "document": document_cfg,
        "editorConfig": {
            "callbackUrl": f"{backend_url}/callback",
            "mode": "view" if is_signed else "edit",
            "lang": "ru",
        },
        "documentType": _get_document_type(ext),
    }

    config["token"] = jwt.encode(config, JWT_SECRET, algorithm="HS256")

    return config
